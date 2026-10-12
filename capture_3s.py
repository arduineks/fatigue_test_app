#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Запись 3 секунд данных с АЦП (STM32) в test/*.xlsx.

Консольный скрипт для быстрой проверки датчика/канала без GUI:

    python capture_3s.py 42        # число 42 попадёт в имя файла
    python capture_3s.py 42 -t 5   # длительность 5 с (по умолчанию 3)

Что делает:
  1. Открывает /dev/ttyACM0 (115200, схема «прочитал — выкинул в USB»).
  2. Синхронизируется на кадры `<BiiffB` (0xAA ... 0xBB, 18 байт).
  3. Пишет все приходящие кадры DURATION секунд.
  4. Сохраняет test/capture_<число>_<дата-время>.xlsx
     (лист «Кадры»: unix-время, время с, RAW, FILTERED, Сила Н,
     Позиция мм; лист «Сводка»: кадров, фактическая частота и т.д.)

Зависимости: pyserial, openpyxl (оба в .venv репозитория).
Порт не передаётся аргументом — по умолчанию /dev/ttyACM0 (можно
переопределить переменной окружения CAPTURE_PORT).
"""

import argparse
import os
import struct
import sys
import time
from datetime import datetime
from pathlib import Path

import serial

# Формат кадра из src/protocol.py (структурно то же; скрипт автономен).
FRAME_FMT = "<BiiffB"
FRAME_SIZE = 18
START = 0xAA
END = 0xBB
BAUD = 115200

REPO_ROOT = Path(__file__).resolve().parent
PORT = os.environ.get("CAPTURE_PORT", "/dev/ttyACM0")


def read_frames(ser, duration_s, on_progress):
    """Читает кадры duration_s секунд; возвращает список строк данных."""
    buf = bytearray()
    rows = []
    t0 = time.time()
    t_next_report = t0 + 0.5

    while True:
        now = time.time()
        if now - t0 >= duration_s:
            break
        if now >= t_next_report:
            on_progress(now - t0, len(rows))
            t_next_report = now + 0.5

        data = ser.read(256)          # блокирующее чтение с timeout
        now = time.time()
        if not data:
            continue
        buf.extend(data)

        # Разбор кадров: ищем 0xAA ... 0xBB (18 байт).
        while len(buf) >= FRAME_SIZE:
            start_idx = buf.find(bytes([START]))
            if start_idx < 0:
                buf.clear()
                break
            if start_idx > 0:
                del buf[:start_idx]
            if len(buf) < FRAME_SIZE:
                break
            if buf[FRAME_SIZE - 1] != END:
                # 0xAA — не начало кадра; сдвиг на 1 байт.
                del buf[:1]
                continue
            try:
                (_, raw, filtered, force_n, current_mm, _) = struct.unpack(
                    FRAME_FMT, bytes(buf[:FRAME_SIZE])
                )
            except struct.error:
                del buf[:1]
                continue
            rows.append((
                time.time(),
                raw,
                filtered,
                force_n,
                current_mm,
            ))
            del buf[:FRAME_SIZE]

    return rows


def save_xlsx(rows, out_path, label):
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Кадры"

    header = [
        "Время(unix)", "Время,с", "RAW", "FILTERED",
        "Сила,Н", "Позиция,мм",
    ]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    t0 = rows[0][0] if rows else 0.0
    for (t, raw, filtered, force_n, current_mm) in rows:
        ws.append([t, round(t - t0, 4), raw, filtered, force_n, current_mm])

    # Сводка.
    ws2 = wb.create_sheet("Сводка")
    dt = rows[-1][0] - rows[0][0] if len(rows) > 1 else 0.0
    summary = [
        ("Метка", label),
        ("Файл создан", datetime.now().isoformat(timespec="seconds")),
        ("Порт", PORT),
        ("Длительность, с", round(dt, 3) if dt else 0),
        ("Кадров записано", len(rows)),
        ("Фактическая частота, Гц", round((len(rows) - 1) / dt, 1) if dt > 0 else "—"),
        ("RAW: min / max", f"{min(r[1] for r in rows)} / {max(r[1] for r in rows)}" if rows else "—"),
        ("FILTERED: min / max", f"{min(r[2] for r in rows)} / {max(r[2] for r in rows)}" if rows else "—"),
        ("Сила Н: min / max", f"{min(r[3] for r in rows):.4f} / {max(r[3] for r in rows):.4f}" if rows else "—"),
        ("Позиция мм: min / max", f"{min(r[4] for r in rows):.3f} / {max(r[4] for r in rows):.3f}" if rows else "—"),
    ]
    for row in summary:
        ws2.append(row)
    for cell in ws2["A"]:
        cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 12
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(
        description="Запись 3 с данных с АЦП в test/*.xlsx"
    )
    parser.add_argument("label", help="число/метка для имени файла")
    parser.add_argument(
        "-t", "--time", type=float, default=3.0,
        help="длительность записи, с (по умолчанию 3)",
    )
    args = parser.parse_args()

    out_dir = REPO_ROOT / "test"
    out_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%d-%m-%Y_%H-%M-%S")
    out_path = out_dir / f"capture_{args.label}_{stamp}.xlsx"

    print(f"Порт: {PORT} ({BAUD})")
    try:
        ser = serial.Serial(PORT, BAUD, timeout=0.02)
    except serial.SerialException as exc:
        sys.exit(f"ОШИБКА: не удалось открыть {PORT}: {exc}")

    print(f"Запись {args.time:.0f} с ...")
    try:
        rows = read_frames(
            ser, args.time,
            on_progress=lambda el, n: print(
                f"  {el:4.1f} с | кадров: {n}", end="\r", flush=True
            ),
        )
    finally:
        ser.close()
    print(f"\n  кадров: {len(rows)}")

    if not rows:
        sys.exit("ОШИБКА: ни одного кадра не получено (устройство не стримит?).")

    dt = rows[-1][0] - rows[0][0]
    rate = (len(rows) - 1) / dt if dt > 0 else 0
    save_xlsx(rows, out_path, args.label)
    print(f"Частота: {rate:.0f} Гц")
    print(f"Сохранено: {out_path}")


if __name__ == "__main__":
    main()
